#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
format_scan 自检：验证「格式语义可查询」这一差异化能力。

构造一张中文表，故意加入：
  - 一整行删除线（模拟"作废"）
  - 一格黄色高亮（模拟"待审阅"）
然后检查 format_scan 能否正确识别、统计、并给出业务含义提示。

这是主流 Excel MCP 完全不具备的能力，必须守住。
"""

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    print(f"  [{'PASS' if cond else 'FAIL'}] {name} {detail if not cond else ''}")
    if cond:
        PASS += 1
    else:
        FAIL += 1


def main() -> int:
    print("=" * 58)
    print(" duduExcel format_scan 自检（格式语义可查询）")
    print("=" * 58)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from duduexcel.advanced import format_scan
    except ImportError as e:
        print(f"\n[SKIP] 依赖缺失：{e}")
        return 0

    tmp = ROOT / "tests" / "_fmt.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws.append(["订单号", "客户", "金额", "状态"])
    ws.append(["A001", "甲公司", 12000, "生效"])
    ws.append(["A002", "乙公司", 8000, "作废"])
    ws.append(["A003", "丙公司", 15000, "待审阅"])

    # 第 3 行整行删除线
    for col in range(1, 5):
        ws.cell(row=3, column=col).font = Font(strike=True)
    # D4 黄色高亮
    ws["D4"].fill = PatternFill("solid", fgColor="FFFF00")
    # 合并 A5:B5（测试 M 语义）
    ws.merge_cells("A5:B5")
    ws["A5"] = "合计"
    wb.save(str(tmp))

    try:
        # 1. 筛选删除线
        r = format_scan(tmp, semantic="S")
        check("能筛出删除线语义", r["hits_count"] == 4, f"got={r['hits_count']}")
        check("删除线命中含 A3/B3/C3/D3",
              sorted(h["cell"] for h in r["hits"]) == ["A3", "B3", "C3", "D3"])
        check("给出业务含义提示（作废）", "作废" in (r["hint"] or ""))

        # 2. 全量统计
        r2 = format_scan(tmp)
        counts = r2["semantic_counts"]
        check("统计含 S 语义", counts.get("S") == 4, f"counts={counts}")
        check("统计含 HL 语义", counts.get("HL") == 1)
        check("统计含 M 语义（合并单元格）", counts.get("M", 0) >= 1)
        check("扫描单元格数合理", r2["cells_scanned"] > 0)

        # 3. 筛选高亮
        r3 = format_scan(tmp, semantic="HL")
        check("能筛出高亮语义", r3["hits_count"] == 1, f"got={r3['hits_count']}")
        check("高亮命中含 D4", any(h["cell"] == "D4" for h in r3["hits"]))
        check("高亮给出业务含义（待审阅）", "待审阅" in (r3["hint"] or ""))

        # 4. 无匹配的语义不应崩
        r4 = format_scan(tmp, semantic="I")
        check("无匹配语义不崩溃", r4["ok"] and r4["hits_count"] == 0)

        # 5. max_cells 截断逻辑
        r5 = format_scan(tmp, semantic="S", max_cells=2)
        check("max_cells 生效", r5["hits_count"] == 2 and r5["truncated"])
    finally:
        tmp.unlink(missing_ok=True)

    print("\n" + "=" * 58)
    print(f" 结果：{PASS} 通过 / {FAIL} 失败")
    print("=" * 58)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
