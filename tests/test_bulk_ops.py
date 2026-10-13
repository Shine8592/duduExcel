#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
P2 广度补齐自检：数据验证 / 去重 / 查找替换

对标 haris-musa/excel-mcp-server（4234★）的高价值能力，
duduExcel 此前缺失，本轮补齐。必须守住，防止回归。

注意：本测试会真实写文件（在 tests/ 下建临时文件后清理）。

修过的两个 openpyxl / Python 坑（测试会覆盖这两个场景）：
1. 有数据验证时 `delete_rows` 会抛
   ValueError: Invalid column index（openpyxl 已知 bug）
   → 改用重建数据法，本测试验证「去重后数据验证仍保留」。

2. 中文列名 `"订单号".isalpha() == True`（Python 中文也是字母）
   → 列字母判断必须加 `isascii()`，本测试用中文列名去重，覆盖该场景。
"""

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    print(f"  [{'PASS' if cond else 'FAIL'}] {name} {detail if not cond else ''}")
    if cond:
        PASS += 1
    else:
        FAIL += 1


def main() -> int:
    print("=" * 58)
    print(" duduExcel P2 自检（数据验证 / 去重 / 查找替换）")
    print("=" * 58)

    try:
        from openpyxl import Workbook, load_workbook
        from duduexcel.advanced import (
            add_data_validation, remove_duplicates, find_replace,
        )
    except ImportError as e:
        print(f"\n[SKIP] 依赖缺失：{e}")
        return 0

    tmp = ROOT / "tests" / "_p2.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws.append(["订单号", "客户", "金额", "状态"])
    ws.append(["A001", "甲公司", 12000, "生效"])
    ws.append(["A002", "乙公司", 8000, "生效"])
    ws.append(["A001", "甲公司", 12000, "生效"])   # 与 A001 重复
    ws.append(["A003", "丙公司", 15000, "待审"])     # "待审" → "待审阅"
    wb.save(str(tmp))

    try:
        # --- 1. 数据验证（下拉）---
        r1 = add_data_validation(tmp, sheet="订单", cell_range="D2:D100",
                                 dv_type="list",
                                 formula1='"生效,作废,待审阅"')
        check("数据验证写入成功", r1.get("ok") is True, str(r1.get("error")))
        check("类型为 list", r1.get("type") == "list")
        check("区域正确", r1.get("cell_range") == "D2:D100")

        # 不支持的类型要给出可用清单
        r1b = add_data_validation(tmp, sheet="订单", cell_range="A1", dv_type="不存在的类型")
        check("非法类型给出可用清单", r1b.get("ok") is False and "可用" in str(r1b.get("error")))

        # --- 2. 去重（中文列名 —— 覆盖 isalpha 坑）---
        r2 = remove_duplicates(tmp, sheet="订单", columns=["订单号"])
        check("去重成功", r2.get("ok") is True, str(r2.get("error")))
        check("删除 1 行重复", r2.get("removed") == 1, f"removed={r2.get('removed')}")
        check("剩余 3 行数据", r2.get("remaining") == 3, f"remaining={r2.get('remaining')}")

        # 关键：去重后数据验证仍保留（覆盖 delete_rows 的 openpyxl bug）
        wb2 = load_workbook(tmp)
        ws2 = wb2["订单"]
        check("去重后数据验证仍保留",
              len(ws2.data_validations.dataValidation) == 1,
              f"got={len(ws2.data_validations.dataValidation)}")

        # 数据正确性：A001/A002/A003，无重复
        rows = [r for r in ws2.iter_rows(values_only=True)]
        check("表头保留", rows[0] == ("订单号", "客户", "金额", "状态"))
        check("去重后无重复订单号", len({r[0] for r in rows[1:]}) == len(rows) - 1)

        wb2.close()

        # --- 3. 查找替换 ---
        r3 = find_replace(tmp, find="待审", replace="待审阅", sheet="订单")
        check("查找替换成功", r3.get("ok") is True)
        check("替换 1 处", r3.get("replaced_count") == 1,
              f"got={r3.get('replaced_count')}")

        # 校验最终值
        wb3 = load_workbook(tmp)
        ws3 = wb3["订单"]
        vals = [r[3] for r in ws3.iter_rows(values_only=True)]
        check("待审已替换为待审阅", "待审阅" in vals)
        check("不再有裸的「待审」", "待审" not in vals)
        wb3.close()

        # --- 4. 边界：find 为空 ---
        r4 = find_replace(tmp, find="", replace="x")
        check("find 为空给出错误", r4.get("ok") is False)

    finally:
        tmp.unlink(missing_ok=True)

    print("\n" + "=" * 58)
    print(f" 结果：{PASS} 通过 / {FAIL} 失败")
    print("=" * 58)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
